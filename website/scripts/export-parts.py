"""Export the public spare-parts list to src/data/parts.json for the website.
Same source and same rules as backend/scripts/make-parts-price-list.py:
docs/Oasis-Parts-List.xlsx, parts with no MRP skipped, display names cleaned.
Run from website/:  python scripts/export-parts.py
"""
import json, re, sys, os
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "..", "docs", "Oasis-Parts-List.xlsx")
OUT = os.path.join(HERE, "..", "src", "data", "parts.json")

# Reuse the display-name fixes from the PDF script so both lists read the same.
sys.path.insert(0, os.path.join(HERE, "..", "..", "oasis-globe", "backend", "scripts"))
CLEAN = {}
src = open(os.path.join(HERE, "..", "..", "oasis-globe", "backend", "scripts", "make-parts-price-list.py"), encoding="utf-8").read()
m = re.search(r"CLEAN = \{(.*?)\n\}", src, re.S)
exec("CLEAN = {" + m.group(1) + "\n}")

def title_case(name):
    if name in CLEAN: return CLEAN[name]
    if name.isupper():
        out = name.title()
        return re.sub(r"\b(Uv|Ro|Uf|Pcb|Smps|Nrv|Ag|Aq|Cc|Tds|Pp|Ea|Nxt|Dt|Kt|Fg)\b", lambda x: x.group(1).upper(), out)
    return name

ws = load_workbook(SRC, data_only=True)["All parts"]
rows = list(ws.iter_rows(values_only=True))
hdr = [str(h) for h in rows[0]]
iN, iB, iM = hdr.index("Part"), hdr.index("Brand"), hdr.index("MRP")
SKIP = {"INSTALLATION CHARGE"}
FITS = {"Oasis": "Most purifiers", "Kent": "Kent", "Aquaguard": "Aquaguard"}
TITLE = {"Oasis": "Oasis spares — fits most purifiers", "Kent": "Kent spares", "Aquaguard": "Aquaguard spares"}
groups = {}
for r in rows[1:]:
    name, brand, mrp = r[iN], r[iB], r[iM]
    if not name or name in SKIP or not mrp or brand in (None, "No brand"): continue
    groups.setdefault(brand, []).append({"part": title_case(str(name).strip()), "fits": FITS.get(brand, brand), "price": int(round(float(mrp)))})
order = [b for b in ("Oasis", "Kent", "Aquaguard") if b in groups] + sorted(b for b in groups if b not in ("Oasis", "Kent", "Aquaguard"))
out = [{"name": TITLE.get(b, f"{b} spares"), "rows": sorted(groups[b], key=lambda x: x["part"].lower())} for b in order]
json.dump(out, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(f"wrote {OUT}: {sum(len(g['rows']) for g in out)} parts in {len(out)} groups")
